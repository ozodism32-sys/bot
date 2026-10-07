"""BAYONNOMA Word hujjatini yaratish (botdan mustaqil ishlaydi).

Foydalanish:
    from docx_builder import build_protocol
    build_protocol(data, ["1.jpg", "2.jpg"], "out.docx")

`data` kalitlari:
    approver_position, approver_name, university, faculty, event_name, event_type,
    number, date (KK.OO.YYYY), city, participants, venue, count, time (SS:DD),
    agenda (list[str]), heard (list[str]), decisions (list[str]),
    secretary, signers (list[{"position": ..., "name": ...}])
"""
from __future__ import annotations

import logging
import tempfile
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from PIL import Image, ImageOps

from utils import date_to_words, dative_suffix, faculty_locative, normalize_uz, parse_date

log = logging.getLogger(__name__)

FONT_NAME = "Times New Roman"
FONT_SIZE = Pt(14)
LINE_SPACING = 1.0

PAGE_W, PAGE_H = Cm(21), Cm(29.7)
MARGIN_LEFT, MARGIN_RIGHT = Cm(3), Cm(1.5)
MARGIN_TOP, MARGIN_BOTTOM = Cm(2), Cm(2)
TEXT_WIDTH_CM = 21 - 3 - 1.5  # 16.5 sm

FIRST_LINE_INDENT = Cm(1.25)
APPROVE_INDENT = Cm(9.0)       # TASDIQLAYMAN bloki o'ng tomonda
IMAGE_HEIGHT_CM = 9.0
IMAGE_CELL_WIDTH_CM = TEXT_WIDTH_CM / 2
IMAGE_MAX_WIDTH_CM = IMAGE_CELL_WIDTH_CM - 0.6
IMAGE_MAX_PX = 1600


# ---------------------------------------------------------------- yordamchilar
def _t(text) -> str:
    return normalize_uz(str(text or "").strip())


def _set_cell_borders_none(table) -> None:
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "nil")
        borders.append(el)
    tbl_pr.append(borders)


def _set_col_widths(table, widths_cm: list[float]) -> None:
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tbl_pr.append(layout)
    for row in table.rows:
        for idx, w in enumerate(widths_cm):
            row.cells[idx].width = Cm(w)
    for idx, col in enumerate(table.columns):
        col.width = Cm(widths_cm[idx])


def _setup_document() -> Document:
    doc = Document()
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.PORTRAIT
    sec.page_width, sec.page_height = PAGE_W, PAGE_H
    sec.left_margin, sec.right_margin = MARGIN_LEFT, MARGIN_RIGHT
    sec.top_margin, sec.bottom_margin = MARGIN_TOP, MARGIN_BOTTOM

    style = doc.styles["Normal"]
    style.font.name = FONT_NAME
    style.font.size = FONT_SIZE
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(attr), FONT_NAME)
    pf = style.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    pf.line_spacing = LINE_SPACING
    return doc


def _para(container, text: str = "", *, bold: bool = False, align=None,
          first_indent=None, left_indent=None, space_before: float = 0,
          space_after: float = 0):
    p = container.add_paragraph()
    if text:
        p.add_run(text).bold = bold
    pf = p.paragraph_format
    if align is not None:
        p.alignment = align
    if first_indent is not None:
        pf.first_line_indent = first_indent
    if left_indent is not None:
        pf.left_indent = left_indent
    pf.space_before = Pt(space_before)
    pf.space_after = Pt(space_after)
    return p


def _cell_para(cell, align=None):
    p = cell.paragraphs[0]
    if align is not None:
        p.alignment = align
    return p


# ---------------------------------------------------------------- bo'limlar
def _approve_block(doc, data: dict) -> None:
    year = ""
    d = parse_date(data.get("date", ""))
    if d:
        year = str(d.year)
    lines = [
        "TASDIQLAYMAN",
        _t(data.get("approver_position")),
        f"___________ {_t(data.get('approver_name'))}",
        f"“____” _________ {year or '______'} yil",
    ]
    for i, line in enumerate(lines):
        _para(doc, line, bold=True, left_indent=APPROVE_INDENT,
              align=WD_ALIGN_PARAGRAPH.LEFT, space_after=6 if i in (1, 2) else 0)


def _title(doc, data: dict) -> None:
    event_type = _t(data.get("event_type")).lower()
    event_name = _t(data.get("event_name"))
    quoted = f"“{event_name}”"
    if event_type == "tashrif":
        quoted += dative_suffix(event_name)
    title = (
        f"{_t(data.get('university'))} {faculty_locative(_t(data.get('faculty')))} "
        f"tahsil olayotgan talaba yoshlari bilan {quoted} {event_type}"
    )
    _para(doc, title, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=18)
    number = _t(data.get("number")) or "____"
    _para(doc, f"B A Y O N N O M A S I № {number}", bold=True,
          align=WD_ALIGN_PARAGRAPH.CENTER, space_after=12)


def _date_line(doc, data: dict) -> None:
    city = _t(data.get("city"))
    if city and not city.lower().endswith("shahri"):
        city = f"{city} shahri"
    p = _para(doc, space_after=12)
    p.paragraph_format.tab_stops.add_tab_stop(Cm(TEXT_WIDTH_CM), WD_TAB_ALIGNMENT.RIGHT)
    p.add_run(f"{date_to_words(data.get('date', ''))}\t{city}").bold = True


def _info_table(doc, data: dict) -> None:
    rows = [
        ("Tadbir ishtirokchilari:", [(_t(data.get("participants")), False)]),
        ("Tadbir o‘tkazish joyi:", [(_t(data.get("venue")), True)]),
        ("Ishtirokchilar soni:", [(f"Jami {data.get('count', '')} nafar", False)]),
        ("Tadbir o‘tkazish vaqti:", [
            (f"{data.get('date', '')}-yil ", False),
            (f"soat {data.get('time', '')}", True),
        ]),
    ]
    table = doc.add_table(rows=len(rows), cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    _set_cell_borders_none(table)
    _set_col_widths(table, [6.0, 0.8, TEXT_WIDTH_CM - 6.8])
    for r, (label, value_runs) in enumerate(rows):
        cells = table.rows[r].cells
        _cell_para(cells[0]).add_run(label).bold = True
        _cell_para(cells[1], WD_ALIGN_PARAGRAPH.CENTER).add_run("–")
        vp = _cell_para(cells[2], WD_ALIGN_PARAGRAPH.LEFT)
        for text, bold in value_runs:
            vp.add_run(text).bold = bold
        for c in cells:
            c.paragraphs[0].paragraph_format.space_after = Pt(4)


def _heading(doc, text: str) -> None:
    p = _para(doc, text, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER,
              space_before=12, space_after=6)
    p.paragraph_format.keep_with_next = True  # sarlavha sahifa oxirida yolg'iz qolmasin


def _numbered(doc, items: list[str]):
    # Raqamlash kod orqali ketma-ket qilinadi: 1, 2, 3 ... (takrorlanmaydi)
    last = None
    for i, item in enumerate((x for x in items if str(x).strip()), start=1):
        last = _para(doc, f"{i}. {_t(item)}", align=WD_ALIGN_PARAGRAPH.JUSTIFY,
                     first_indent=FIRST_LINE_INDENT)
    return last


def _paragraphs(doc, items: list[str]) -> None:
    for item in items:
        if str(item).strip():
            _para(doc, _t(item), align=WD_ALIGN_PARAGRAPH.JUSTIFY,
                  first_indent=FIRST_LINE_INDENT)


def _prepare_image(src: str, tmp_dir: Path, idx: int) -> tuple[str, int, int] | None:
    """Rasmni burilishini to'g'rilaydi, siqadi va JPEG qiladi."""
    try:
        with Image.open(src) as im:
            im = ImageOps.exif_transpose(im)
            if im.mode not in ("RGB", "L"):
                bg = Image.new("RGB", im.size, (255, 255, 255))
                rgba = im.convert("RGBA")
                bg.paste(rgba, mask=rgba.split()[-1])
                im = bg
            elif im.mode == "L":
                im = im.convert("RGB")
            im.thumbnail((IMAGE_MAX_PX, IMAGE_MAX_PX), Image.LANCZOS)
            out = tmp_dir / f"img_{idx}.jpg"
            im.save(out, "JPEG", quality=85, optimize=True)
            return str(out), im.width, im.height
    except Exception:
        log.exception("Rasmni qayta ishlab bo'lmadi: %s", src)
        return None


def _images(doc, images: list[str], tmp_dir: Path) -> None:
    prepared = [p for i, src in enumerate(images[:6]) if (p := _prepare_image(src, tmp_dir, i))]
    if not prepared:
        return
    _para(doc, space_after=6)
    for start in range(0, len(prepared), 2):
        pair = prepared[start:start + 2]
        if len(pair) == 2:
            # Ikkala rasm bir xil balandlikda, katakka sig'adigan qilib
            height = min(IMAGE_HEIGHT_CM, *(IMAGE_MAX_WIDTH_CM * h / w for _, w, h in pair))
            table = doc.add_table(rows=1, cols=2)
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            _set_cell_borders_none(table)
            _set_col_widths(table, [IMAGE_CELL_WIDTH_CM, IMAGE_CELL_WIDTH_CM])
            for cell, (path, _, _) in zip(table.rows[0].cells, pair):
                p = _cell_para(cell, WD_ALIGN_PARAGRAPH.CENTER)
                p.add_run().add_picture(path, height=Cm(height))
        else:
            path, w, h = pair[0]
            height = min(IMAGE_HEIGHT_CM, 12.0 * h / w)
            p = _para(doc, align=WD_ALIGN_PARAGRAPH.CENTER)
            p.add_run().add_picture(path, height=Cm(height))
        _para(doc, space_after=4)


def _signatures(doc, data: dict) -> None:
    signers = [s for s in data.get("signers") or [] if s.get("name")]
    lines = [f"{_t(s.get('position'))}: {_t(s['name'])}" for s in signers]
    lines.append(f"Yig‘ilish kotibi: {_t(data.get('secretary'))}")
    for i, line in enumerate(lines):
        p = _para(doc, line, bold=True, align=WD_ALIGN_PARAGRAPH.RIGHT,
                  space_before=24 if i == 0 else 12)
        # imzolar bir-biridan va oxirgi qarordan ajralib qolmasin
        p.paragraph_format.keep_with_next = i < len(lines) - 1


# ---------------------------------------------------------------- asosiy
def build_protocol(data: dict, images: list[str], out_path: str) -> str:
    """Bayonnoma DOCX faylini yaratadi va yo'lini qaytaradi."""
    out = Path(out_path)
    out.parent.mkdir(parents=True, exist_ok=True)
    doc = _setup_document()

    _approve_block(doc, data)
    _title(doc, data)
    _date_line(doc, data)
    _info_table(doc, data)

    _heading(doc, "KUN TARTIBI:")
    _numbered(doc, data.get("agenda") or [])

    _heading(doc, "ESHITILDI:")
    _paragraphs(doc, data.get("heard") or [])

    with tempfile.TemporaryDirectory() as tmp:
        _images(doc, images or [], Path(tmp))

        _heading(doc, "QAROR QILINDI:")
        last_decision = _numbered(doc, data.get("decisions") or [])
        if last_decision is not None:
            last_decision.paragraph_format.keep_with_next = True

        _signatures(doc, data)
        doc.save(out)
    log.info("DOCX yaratildi: %s", out)
    return str(out)
